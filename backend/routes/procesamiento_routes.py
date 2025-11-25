from fastapi import APIRouter, UploadFile, File, Form, Depends, HTTPException, Query
from fastapi.responses import JSONResponse, StreamingResponse
from sqlalchemy.orm import Session
from controllers.procesamiento_controller import ProcesamientoController
from services.auth_service import get_current_user
from services.modelo_service import ModeloService
from models.user import User
from models.empresa import Empresa
from config.database import get_db
from utils.number_utils import remove_trailing_zeros
import logging
import pandas as pd
import io
from datetime import datetime

router = APIRouter()
logger = logging.getLogger(__name__)

@router.post("/procesar-excel")
async def procesar_excel(
    file: UploadFile = File(...),
    nit_empresa: str = Form("901906032"),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):

    try:
        resultado = await ProcesamientoController.ejecutar_script_procesamiento(
            file=file,
            nit_empresa=nit_empresa,
            user_id=current_user.id,  # Pasar user_id del usuario autenticado
            db=db
        )
        return JSONResponse(content={"success": True, "data": resultado})
    except Exception as e:
        logger.error(f"Error procesando archivo: {str(e)}")
        return JSONResponse(
            status_code=500,
            content={"success": False, "error": str(e)}
        )

@router.get("/archivos-procesados")
async def listar_archivos_procesados(
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    
    try:
        archivos = ProcesamientoController.listar_archivos_procesados(db, current_user.id)  # Pasar user_id
        return JSONResponse(
            content={
                "success": True,
                "data": archivos,
                "total": len(archivos)
            }
        )
    except Exception as e:
        logger.error(f"Error listando archivos: {str(e)}")
        return JSONResponse(
            status_code=500,
            content={
                "success": False,
                "error": f"Error listando archivos: {str(e)}"
            }
        )

@router.get("/descargar-archivo/{archivo_id}")
async def descargar_archivo(
    archivo_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):

    try:
        archivo = await ProcesamientoController.descargar_archivo(archivo_id, current_user.id, db)  # Pasar user_id
        if archivo:
            return JSONResponse(
                content={
                    "success": True,
                    "data": archivo.to_dict()
                }
            )
        return JSONResponse(
            status_code=404,
            content={
                "success": False,
                "error": "Archivo no encontrado"
            }
        )
    except Exception as e:
        logger.error(f"Error descargando archivo: {str(e)}")
        return JSONResponse(
            status_code=500,
            content={
                "success": False,
                "error": f"Error descargando archivo: {str(e)}"
            }
        )

@router.delete("/eliminar-archivo/{archivo_id}")
async def eliminar_archivo(
    archivo_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):

    try:
        resultado = await ProcesamientoController.eliminar_archivo(archivo_id, current_user.id, db)  # Pasar user_id
        if resultado:
            return JSONResponse(
                content={
                    "success": True,
                    "message": "Archivo eliminado correctamente"
                }
            )
        return JSONResponse(
            status_code=404,
            content={
                "success": False,
                "error": "Archivo no encontrado"
            }
        )
    except Exception as e:
        logger.error(f"Error eliminando archivo: {str(e)}")
        return JSONResponse(
            status_code=500,
            content={
                "success": False,
                "error": f"Error eliminando archivo: {str(e)}"
            }
        )

@router.post("/generar-modelo-ventas/{archivo_id}")
async def generar_modelo_ventas(
    archivo_id: int,
    nit_empresa: str = Form("901906032"),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):

    try:
        resultado = ModeloService.generar_modelo_ventas(
            archivo_id=archivo_id,
            user_id=current_user.id,
            nit_empresa=nit_empresa,
            db=db
        )
        return JSONResponse(content={"success": True, "data": resultado})
    except Exception as e:
        logger.error(f"Error generando modelo de ventas: {str(e)}")
        return JSONResponse(
            status_code=500,
            content={"success": False, "error": str(e)}
        )

@router.post("/generar-modelo-compras/{archivo_id}")
async def generar_modelo_compras(
    archivo_id: int,
    nit_empresa: str = Form("901906032"),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    
    try:
        resultado = ModeloService.generar_modelo_compras(
            archivo_id=archivo_id,
            user_id=current_user.id,
            nit_empresa=nit_empresa,
            db=db
        )
        return JSONResponse(content={"success": True, "data": resultado})
    except Exception as e:
        logger.error(f"Error generando modelo de compras: {str(e)}")
        return JSONResponse(
            status_code=500,
            content={"success": False, "error": str(e)}
        )

@router.get("/detectar-dian/{archivo_id}")
async def detectar_archivo_dian(
    archivo_id: int,
    nit_empresa: str = "901906032",
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):

    try:
        resultado = ModeloService.detectar_y_procesar_dian(
            archivo_id=archivo_id,
            user_id=current_user.id,
            nit_empresa=nit_empresa,
            db=db
        )
        return JSONResponse(content={"success": True, "data": resultado})
    except Exception as e:
        logger.error(f"Error detectando archivo DIAN: {str(e)}")
        return JSONResponse(
            status_code=500,
            content={"success": False, "error": str(e)}
        )

# ===== NUEVOS ENDPOINTS PARA PLANTILLAS EXCEL =====

@router.post("/procesar-terceros")
async def procesar_terceros(
    file: UploadFile = File(...),
    nit_empresa: str = Form("901906032"),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):

    try:
        resultado = await ProcesamientoController.ejecutar_procesamiento_terceros(
            file=file,
            nit_empresa=nit_empresa,
            user_id=current_user.id,
            db=db
        )
        return JSONResponse(content={"success": True, "data": resultado})
    except Exception as e:
        logger.error(f"Error procesando archivo de terceros: {str(e)}")
        return JSONResponse(
            status_code=500,
            content={"success": False, "error": str(e)}
        )

@router.get("/descargar-plantilla-cuentas")
async def descargar_plantilla_cuentas(
    empresa_id: int = Query(..., description="ID de la empresa"),
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):

    try:
        # Verificar que la empresa existe y pertenece al usuario
        from models.empresa import Empresa
        from models.cuenta_importada import CuentaImportada
        from models.base import Base
        
        # Si empresa_id es 0, es una empresa nueva - usar plantilla por defecto con ejemplos
        if empresa_id == 0:
            empresa_nombre = "Nueva_Empresa"
            cuentas_empresa = []  # Forzar plantilla con ejemplos para empresa nueva
        else:
            empresa = db.query(Empresa).filter(
                Empresa.id == empresa_id,
                Empresa.usuario_id == current_user.id,
                Empresa.is_active == True
            ).first()
            
            if not empresa:
                raise HTTPException(status_code=404, detail="Empresa no encontrada")
            
            empresa_nombre = empresa.razon_social
        
        try:
            # Intentar crear la tabla si no existe (usando ORM)
            CuentaImportada.__table__.create(bind=db.get_bind(), checkfirst=True)
            logger.info("📊 Tabla cuentas_importadas verificada/creada")
        except Exception as table_error:
            logger.warning(f"⚠️ Error verificando tabla: {table_error}")
        
        # Buscar cuentas específicas de esta empresa (solo si empresa_id > 0)
        if empresa_id > 0:
            try:
                # Buscar cuentas de la empresa específica
                cuentas_empresa = db.query(CuentaImportada).filter(
                    CuentaImportada.user_id == current_user.id,
                    CuentaImportada.empresa_id == empresa_id
                ).order_by(CuentaImportada.fecha_actualizacion.desc()).all()
                logger.info(f"📊 Cuentas encontradas para empresa {empresa_nombre} (ID: {empresa_id}): {len(cuentas_empresa)}")
            except Exception as db_error:
                logger.error(f"❌ Error consultando BD: {str(db_error)}")
                cuentas_empresa = []
        # Si empresa_id = 0, NO buscar cuentas, usar plantilla con ejemplos
        
        if cuentas_empresa:
            plantilla_data = {
                "NIT": [remove_trailing_zeros(cuenta.nit) for cuenta in cuentas_empresa],
                "NOMBRE": [cuenta.nombre for cuenta in cuentas_empresa],
                "CUENTA": [remove_trailing_zeros(cuenta.cuenta) for cuenta in cuentas_empresa],
                "CUENTA_EXENTA": [remove_trailing_zeros(cuenta.cuenta_exenta) if cuenta.cuenta_exenta else "" for cuenta in cuentas_empresa],
                "CUENTA_IVA": [remove_trailing_zeros(cuenta.cuenta_iva) if cuenta.cuenta_iva else "" for cuenta in cuentas_empresa],
                "CUENTA_CONTRAPARTIDA": [remove_trailing_zeros(cuenta.cuenta_contrapartida) if cuenta.cuenta_contrapartida else "" for cuenta in cuentas_empresa]
            }
            tipo_archivo = "Datos"
            logger.info(f"Generando Excel con {len(cuentas_empresa)} cuentas de {empresa_nombre}")
        else:
            plantilla_data = {
                "NIT": ["900123456", "800654321", "900555444"],
                "NOMBRE": ["Ejemplo Proveedor A - Todas las cuentas", "Ejemplo Proveedor B - Solo algunas", "Ejemplo Proveedor C - Solo base"],
                "CUENTA": ["51050101", "51351001", "51352001"],
                "CUENTA_EXENTA": ["51050102", "", ""],
                "CUENTA_IVA": ["24080101", "24080501", ""],
                "CUENTA_CONTRAPARTIDA": ["22050101", "", ""]
            }
            tipo_archivo = "Plantilla"
            logger.info(f"Generando plantilla por defecto para empresa {empresa_nombre}")
        
        df = pd.DataFrame(plantilla_data, dtype=str)
        
        columnas_numericas = ["NIT", "CUENTA", "CUENTA_EXENTA", "CUENTA_IVA", "CUENTA_CONTRAPARTIDA"]
        for col in columnas_numericas:
            if col in df.columns:
                df[col] = df[col].apply(remove_trailing_zeros)
                df[col] = df[col].astype(str)
        
        buffer = io.BytesIO()
        
        with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
            df.to_excel(writer, index=False, sheet_name="Cuentas_Compras")
            
            from openpyxl.styles import Font, PatternFill, Alignment
            workbook = writer.book
            worksheet = writer.sheets["Cuentas_Compras"]
            
            for row in worksheet.iter_rows(min_row=2, max_row=worksheet.max_row):
                if row[0].value:
                    row[0].number_format = '@'
                if len(row) > 2 and row[2].value:
                    row[2].number_format = '@'
                if len(row) > 3 and row[3].value:
                    row[3].number_format = '@'
                if len(row) > 4 and row[4].value:
                    row[4].number_format = '@'
                if len(row) > 5 and row[5].value:
                    row[5].number_format = '@'
            
            # Agregar formato a los headers
            header_font = Font(bold=True, color="FFFFFF")
            header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
            header_alignment = Alignment(horizontal="center", vertical="center")
            
            # Aplicar formato a los headers
            for cell in worksheet[1]:
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = header_alignment
            
            # Ajustar ancho de columnas
            worksheet.column_dimensions["A"].width = 15  # NIT
            worksheet.column_dimensions["B"].width = 40  # NOMBRE
            worksheet.column_dimensions["C"].width = 18  # CUENTA
            worksheet.column_dimensions["D"].width = 18  # CUENTA_EXENTA
            worksheet.column_dimensions["E"].width = 18  # CUENTA_IVA
            worksheet.column_dimensions["F"].width = 25  # CUENTA_CONTRAPARTIDA
        
        buffer.seek(0)
        
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        empresa_nombre_archivo = empresa_nombre.replace(" ", "_").replace(".", "")
        filename = f"{tipo_archivo}_Cuentas_{empresa_nombre_archivo}_{timestamp}.xlsx"
        
        logger.info(f"📝 {tipo_archivo} generada exitosamente para {empresa_nombre}: {filename}")
        
        return StreamingResponse(
            io.BytesIO(buffer.getvalue()),
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f"attachment; filename={filename}"}
        )
        
    except Exception as e:
        logger.error(f"❌ Error generando plantilla: {str(e)}")
        return JSONResponse(
            status_code=500,
            content={"success": False, "error": f"Error generando plantilla: {str(e)}"}
        )

@router.post("/importar-archivo-cuentas")
async def importar_archivo_cuentas(
    file: UploadFile = File(...),
    empresa_id: int = Form(...),  # Requerido - cada empresa debe tener sus propias cuentas
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):

    try:
        # Verificar que sea un archivo Excel
        if not file.filename.lower().endswith((".xlsx", ".xls")):
            return JSONResponse(
                status_code=400,
                content={"success": False, "error": "El archivo debe ser de formato Excel (.xlsx o .xls)"}
            )
        
        # Leer el contenido del archivo
        content = await file.read()
        
        # Cargar el DataFrame
        df = pd.read_excel(io.BytesIO(content))
        
        # Validar que existan las columnas REQUERIDAS
        campos_requeridos = ["NIT", "NOMBRE", "CUENTA"]
        
        # Campos OPCIONALES
        campos_opcionales = ["CUENTA_EXENTA", "CUENTA_IVA", "CUENTA_CONTRAPARTIDA"]
        
        # Mapear columnas flexiblemente
        mapeo_columnas = {}
        
        # Buscar columnas requeridas
        for campo in campos_requeridos:
            columna_encontrada = None
            for col_archivo in df.columns:
                col_normalizada = col_archivo.upper().strip()
                if col_normalizada == campo or col_normalizada in [campo, campo.lower(), campo.capitalize()]:
                    columna_encontrada = col_archivo
                    break
            
            if columna_encontrada is None:
                return JSONResponse(
                    status_code=400,
                    content={
                        "success": False, 
                        "error": f"El archivo debe contener la columna '{campo}'. Columnas encontradas: {list(df.columns)}"
                    }
                )
            mapeo_columnas[campo] = columna_encontrada
        
        # Buscar columnas opcionales (no dar error si no existen)
        for campo in campos_opcionales:
            columna_encontrada = None
            for col_archivo in df.columns:
                col_normalizada = col_archivo.upper().strip()
                if col_normalizada == campo or col_normalizada in [campo, campo.lower(), campo.capitalize()]:
                    columna_encontrada = col_archivo
                    break
            if columna_encontrada:
                mapeo_columnas[campo] = columna_encontrada
                logger.info(f"✅ Columna opcional '{campo}' encontrada")
            else:
                logger.info(f"ℹ️ Columna opcional '{campo}' no encontrada (se usará configuración empresa)")
        
        # Normalizar los nombres de las columnas
        df_normalizado = pd.DataFrame()
        df_normalizado["NIT"] = df[mapeo_columnas["NIT"]].astype(str)
        df_normalizado["NOMBRE"] = df[mapeo_columnas["NOMBRE"]].astype(str)
        df_normalizado["CUENTA"] = df[mapeo_columnas["CUENTA"]].astype(str)
        
        # Agregar columnas opcionales si existen
        if "CUENTA_EXENTA" in mapeo_columnas:
            df_normalizado["CUENTA_EXENTA"] = df[mapeo_columnas["CUENTA_EXENTA"]].astype(str)
        else:
            df_normalizado["CUENTA_EXENTA"] = ""  # Vacío si no existe
        
        if "CUENTA_IVA" in mapeo_columnas:
            df_normalizado["CUENTA_IVA"] = df[mapeo_columnas["CUENTA_IVA"]].astype(str)
        else:
            df_normalizado["CUENTA_IVA"] = ""  # Vacío si no existe
        
        if "CUENTA_CONTRAPARTIDA" in mapeo_columnas:
            df_normalizado["CUENTA_CONTRAPARTIDA"] = df[mapeo_columnas["CUENTA_CONTRAPARTIDA"]].astype(str)
        else:
            df_normalizado["CUENTA_CONTRAPARTIDA"] = ""  # Vacío si no existe
        
        # Validar que no haya filas vacías en campos críticos
        filas_vacias = df_normalizado[(df_normalizado["NIT"].str.strip() == "") | 
                                    (df_normalizado["NIT"].isna()) |
                                    (df_normalizado["CUENTA"].str.strip() == "") | 
                                    (df_normalizado["CUENTA"].isna())]
        
        if not filas_vacias.empty:
            logger.warning(f"⚠️ Se encontraron {len(filas_vacias)} filas con NIT o CUENTA vacía. Se eliminarán.")
            df_normalizado = df_normalizado.dropna(subset=["NIT", "CUENTA"])
            df_normalizado = df_normalizado[(df_normalizado["NIT"].str.strip() != "") &
            (df_normalizado["CUENTA"].str.strip() != "")]
        
        df_normalizado["NIT"] = df_normalizado["NIT"].str.strip()
        df_normalizado["NOMBRE"] = df_normalizado["NOMBRE"].str.strip().fillna("")
        df_normalizado["CUENTA"] = df_normalizado["CUENTA"].str.strip()
        
        columnas_numericas = ["NIT", "CUENTA", "CUENTA_EXENTA", "CUENTA_IVA", "CUENTA_CONTRAPARTIDA"]
        for col in columnas_numericas:
            if col in df_normalizado.columns:
                df_normalizado[col] = df_normalizado[col].apply(remove_trailing_zeros)
        
        logger.info(f"Limpieza aplicada - Total registros: {len(df_normalizado)}")
        
        # GUARDAR EN BASE DE DATOS
        from models.cuenta_importada import CuentaImportada
        
        # Asegurar que la tabla existe (crear si no existe)
        try:
            CuentaImportada.__table__.create(bind=db.get_bind(), checkfirst=True)
            logger.info("📊 Tabla cuentas_importadas verificada/creada")
        except Exception as table_error:
            logger.warning(f"⚠️ Error verificando tabla: {table_error}")
        
        # Obtener NITs del Excel para saber cuáles mantener
        nits_en_excel = set(df_normalizado["NIT"].unique())
        logger.info(f"📊 NITs en el Excel: {len(nits_en_excel)}")
        
        # Validar que empresa_id sea válido (debe ser > 0, no puede ser None o 0)
        if not empresa_id or empresa_id <= 0:
            return JSONResponse(
                status_code=400,
                content={"success": False, "error": "Se requiere un empresa_id válido. Las cuentas deben estar asociadas a una empresa específica."}
            )
        
        # Verificar que la empresa existe y pertenece al usuario
        empresa = db.query(Empresa).filter(
            Empresa.id == empresa_id,
            Empresa.usuario_id == current_user.id,
            Empresa.is_active == True
        ).first()
        
        if not empresa:
            return JSONResponse(
                status_code=404,
                content={"success": False, "error": "Empresa no encontrada o no tienes permisos para acceder a ella."}
            )
        
        # Construir filtros para identificar cuentas existentes de esta empresa/usuario
        filtros_existentes = [
            CuentaImportada.user_id == current_user.id,
            CuentaImportada.empresa_id == empresa_id
        ]
        
        # Contar cuentas antes de procesar
        total_antes = db.query(CuentaImportada).filter(*filtros_existentes).count()
        logger.info(f"📊 Total de cuentas antes de importar: {total_antes}")
        
        registros_guardados = 0
        registros_actualizados = 0
        
        for _, fila in df_normalizado.iterrows():
            nit = fila["NIT"]
            nombre = fila["NOMBRE"]
            cuenta = fila["CUENTA"]
            cuenta_exenta = fila.get("CUENTA_EXENTA", "").strip() or None  # None si está vacío
            cuenta_iva = fila.get("CUENTA_IVA", "").strip() or None  # None si está vacío
            cuenta_contrapartida = fila.get("CUENTA_CONTRAPARTIDA", "").strip() or None  # None si está vacío
            
            # Verificar si ya existe un registro con el mismo NIT para este usuario y empresa
            filtros = [
                CuentaImportada.user_id == current_user.id,
                CuentaImportada.empresa_id == empresa_id,
                CuentaImportada.nit == nit
            ]
            
            cuenta_existente = db.query(CuentaImportada).filter(*filtros).first()
            
            if cuenta_existente:
                # Actualizar registro existente
                cuenta_existente.nombre = nombre
                cuenta_existente.cuenta = cuenta
                cuenta_existente.cuenta_exenta = cuenta_exenta
                cuenta_existente.cuenta_iva = cuenta_iva
                cuenta_existente.cuenta_contrapartida = cuenta_contrapartida
                cuenta_existente.fecha_actualizacion = datetime.now()
                registros_actualizados += 1
                logger.info(f"🔄 Actualizando cuenta existente: NIT {nit} - Gravada: {cuenta}, Exenta: {cuenta_exenta or 'N/A'}, IVA: {cuenta_iva or 'N/A'}, Contrapartida: {cuenta_contrapartida or 'config'}")
            else:
                # Crear nuevo registro
                nueva_cuenta = CuentaImportada(
                    user_id=current_user.id,
                    empresa_id=empresa_id,
                    nit=nit,
                    nombre=nombre,
                    cuenta=cuenta,
                    cuenta_exenta=cuenta_exenta,
                    cuenta_iva=cuenta_iva,
                    cuenta_contrapartida=cuenta_contrapartida
                )
                db.add(nueva_cuenta)
                registros_guardados += 1
                logger.info(f"➕ Creando nueva cuenta: NIT {nit} - Gravada: {cuenta}, Exenta: {cuenta_exenta or 'N/A'}, IVA: {cuenta_iva or 'N/A'}, Contrapartida: {cuenta_contrapartida or 'config'}")
        
        # Eliminar cuentas que estaban antes pero NO están en el Excel nuevo
        # Esto asegura que el Excel importado reemplace completamente las cuentas anteriores
        cuentas_eliminadas = 0
        if nits_en_excel:
            # Obtener todas las cuentas existentes
            cuentas_existentes = db.query(CuentaImportada).filter(*filtros_existentes).all()
            
            for cuenta_existente in cuentas_existentes:
                # Si el NIT de la cuenta existente NO está en el Excel nuevo, eliminarla
                if cuenta_existente.nit not in nits_en_excel:
                    db.delete(cuenta_existente)
                    cuentas_eliminadas += 1
                    logger.info(f"🗑️ Eliminando cuenta que no está en el Excel nuevo: NIT {cuenta_existente.nit}")
        
        # Confirmar cambios en la base de datos
        db.commit()
        
        # Obtener el total real de cuentas después de la importación
        total_despues = db.query(CuentaImportada).filter(*filtros_existentes).count()
        logger.info(f"📊 Total de cuentas después de importar: {total_despues}")
        
        # Obtener la fecha de última actualización más reciente
        from sqlalchemy import func
        ultima_fecha = db.query(func.max(CuentaImportada.fecha_actualizacion)).filter(*filtros_existentes).scalar()
        if not ultima_fecha:
            ultima_fecha = db.query(func.max(CuentaImportada.fecha_importacion)).filter(*filtros_existentes).scalar()
        
        total_procesados = registros_guardados + registros_actualizados
        mensaje = f"Archivo procesado exitosamente. {registros_guardados} registros nuevos, {registros_actualizados} actualizados, {cuentas_eliminadas} eliminados. Total en Excel: {total_procesados}, Total en sistema: {total_despues}"
        
        logger.info(f"✅ {mensaje}")
        
        return JSONResponse(
            content={
                "success": True,
                "data": {
                    "valido": True,
                    "registros_nuevos": registros_guardados,
                    "registros_actualizados": registros_actualizados,
                    "registros_eliminados": cuentas_eliminadas,
                    "total_procesados": total_procesados,
                    "total_cuentas_sistema": total_despues,  # Total real de cuentas en el sistema
                    "ultima_actualizacion": ultima_fecha.isoformat() if ultima_fecha else None,
                    "mensaje": mensaje
                }
            }
        )
        
    except Exception as e:
        db.rollback()
        logger.error(f"❌ Error validando archivo: {str(e)}")
        return JSONResponse(
            status_code=500,
            content={"success": False, "error": f"Error procesando archivo: {str(e)}"}
        )

@router.post("/test-upload")
async def test_upload(
    file: UploadFile = File(...),
    current_user: User = Depends(get_current_user)
):

    try:
        logger.info(f"📁 Archivo recibido: {file.filename}")
        logger.info(f"📁 Tipo de contenido: {file.content_type}")
        logger.info(f"📁 Tamaño: {file.size if hasattr(file, 'size') else 'Desconocido'}")
        
        content = await file.read()
        logger.info(f"📁 Contenido leído: {len(content)} bytes")
        
        return JSONResponse(
            content={
                "success": True,
                "message": "Archivo recibido correctamente",
                "filename": file.filename,
                "content_type": file.content_type,
                "size": len(content)
            }
        )
    except Exception as e:
        logger.error(f"❌ Error en test upload: {str(e)}")
        return JSONResponse(
            status_code=500,
            content={"success": False, "error": str(e)}
        )

@router.get("/estado-cuentas-importadas")
async def obtener_estado_cuentas_importadas(
    empresa_id: int = None,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):

    try:
        from models.cuenta_importada import CuentaImportada
        
        # Asegurar que la tabla existe
        try:
            CuentaImportada.__table__.create(bind=db.get_bind(), checkfirst=True)
        except Exception:
            pass
        
        # Obtener estadísticas
        filtros = [CuentaImportada.user_id == current_user.id]
        
        if empresa_id and empresa_id > 0:
            # Filtrar por empresa específica
            filtros.append(CuentaImportada.empresa_id == empresa_id)
        elif empresa_id == 0:
            # Para empresa nueva, no mostrar cuentas
            return JSONResponse(
                content={
                    "success": True,
                    "data": {
                        "tiene_cuentas": False,
                        "total_cuentas": 0,
                        "ultima_actualizacion": None,
                        "mensaje": "Descarga la plantilla Excel para comenzar"
                    }
                }
            )
        else:
            # Si no se proporciona empresa_id, no mostrar cuentas
            return JSONResponse(
                content={
                    "success": True,
                    "data": {
                        "tiene_cuentas": False,
                        "total_cuentas": 0,
                        "ultima_actualizacion": None,
                        "mensaje": "Descarga la plantilla Excel para comenzar"
                    }
                }
            )
        
        # Obtener el total real de cuentas
        total_cuentas = db.query(CuentaImportada).filter(*filtros).count()
        
        if total_cuentas > 0:
            # Obtener la fecha de última actualización más reciente de todas las cuentas
            # Esto asegura que siempre muestre la fecha real de la última importación
            from sqlalchemy import func
            ultima_fecha = db.query(func.max(CuentaImportada.fecha_actualizacion)).filter(*filtros).scalar()
            
            # Si no hay fecha de actualización, usar fecha de importación
            if not ultima_fecha:
                ultima_fecha = db.query(func.max(CuentaImportada.fecha_importacion)).filter(*filtros).scalar()
            
            return JSONResponse(
                content={
                    "success": True,
                    "data": {
                        "tiene_cuentas": True,
                        "total_cuentas": total_cuentas,  # Total real de cuentas en el sistema
                        "ultima_actualizacion": ultima_fecha.isoformat() if ultima_fecha else None,
                        "mensaje": f"Tienes {total_cuentas} cuentas importadas"
                    }
                }
            )
        else:
            return JSONResponse(
                content={
                    "success": True,
                    "data": {
                        "tiene_cuentas": False,
                        "total_cuentas": 0,
                        "ultima_actualizacion": None,
                        "mensaje": "No hay cuentas importadas"
                    }
                }
            )
    except Exception as e:
        logger.error(f"❌ Error obteniendo estado: {str(e)}")
        return JSONResponse(
            status_code=500,
            content={"success": False, "error": f"Error obteniendo estado: {str(e)}"}
        )

@router.get("/listar-cuentas-importadas")
async def listar_cuentas_importadas(
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):

    try:
        from models.cuenta_importada import CuentaImportada
        
        # Asegurar que la tabla existe
        try:
            CuentaImportada.__table__.create(bind=db.get_bind(), checkfirst=True)
        except Exception:
            pass
        
        # Obtener todas las cuentas del usuario
        cuentas = db.query(CuentaImportada).filter(
            CuentaImportada.user_id == current_user.id
        ).order_by(CuentaImportada.nit).all()
        
        # Convertir a diccionarios
        cuentas_data = [cuenta.to_dict() for cuenta in cuentas]
        
        return JSONResponse(
            content={
                "success": True,
                "data": cuentas_data,
                "total": len(cuentas_data)
            }
        )
    except Exception as e:
        logger.error(f"❌ Error listando cuentas: {str(e)}")
        return JSONResponse(
            status_code=500,
            content={"success": False, "error": f"Error listando cuentas: {str(e)}"}
        )

@router.delete("/eliminar-cuenta-importada/{nit}")
async def eliminar_cuenta_importada(
    nit: str,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):

    try:
        from models.cuenta_importada import CuentaImportada
        
        # Asegurar que la tabla existe
        try:
            CuentaImportada.__table__.create(bind=db.get_bind(), checkfirst=True)
        except Exception:
            pass
        
        # Buscar y eliminar la cuenta
        cuenta = db.query(CuentaImportada).filter(
            CuentaImportada.user_id == current_user.id,
            CuentaImportada.nit == nit
        ).first()
        
        if not cuenta:
            return JSONResponse(
                status_code=404,
                content={"success": False, "error": f"No se encontró cuenta con NIT {nit}"}
            )
        
        db.delete(cuenta)
        db.commit()
        
        return JSONResponse(
            content={
                "success": True,
                "message": f"Cuenta con NIT {nit} eliminada correctamente"
            }
        )
    except Exception as e:
        logger.error(f"❌ Error eliminando cuenta: {str(e)}")
        db.rollback()
        return JSONResponse(
            status_code=500,
            content={"success": False, "error": f"Error eliminando cuenta: {str(e)}"}
        )

@router.delete("/limpiar-todas-cuentas-importadas")
async def limpiar_todas_cuentas_importadas(
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):

    try:
        from models.cuenta_importada import CuentaImportada
        
        # Asegurar que la tabla existe
        try:
            CuentaImportada.__table__.create(bind=db.get_bind(), checkfirst=True)
        except Exception:
            pass
        
        # Contar cuentas antes de eliminar
        total_cuentas = db.query(CuentaImportada).filter(
            CuentaImportada.user_id == current_user.id
        ).count()
        
        if total_cuentas == 0:
            return JSONResponse(
                content={
                    "success": True,
                    "message": "No hay cuentas importadas para eliminar",
                    "cuentas_eliminadas": 0
                }
            )
        
        # Eliminar todas las cuentas del usuario
        db.query(CuentaImportada).filter(
            CuentaImportada.user_id == current_user.id
        ).delete()
        
        db.commit()
        
        logger.info(f"🗑️ Usuario {current_user.id} eliminó {total_cuentas} cuentas importadas")
        
        return JSONResponse(
            content={
                "success": True,
                "message": f"Se eliminaron {total_cuentas} cuentas importadas correctamente",
                "cuentas_eliminadas": total_cuentas
            }
        )
    except Exception as e:
        logger.error(f"❌ Error eliminando todas las cuentas: {str(e)}")
        db.rollback()
        return JSONResponse(
            status_code=500,
            content={"success": False, "error": f"Error eliminando cuentas: {str(e)}"}
        )